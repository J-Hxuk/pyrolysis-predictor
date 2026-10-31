"""
test_prediction_files.py - Task 9 완료 기준 검증.

합성 자료로만 테스트한다. 실제 연구 변수·단위는 사용하지 않는다.

확인 항목:
1. 마지막 행 오류도 찾아 파일 전체를 실패 처리한다 (candidates=None).
2. 표시할 오류 수를 제한해도 전체 검사·전체 오류 수 집계는 수행한다.
3. 잘못된 문자열을 결측으로 바꾸어 통과시키지 않는다.
4. 중복 헤더·미지 열·필수 열 누락·빈 파일·비유한 값·혼합/불일치 규격 ID 처리.
5. 열 순서와 무관하게 정의된 ID 로 읽는다.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from pyrolysis.config import build_input_schema, load_problem
from pyrolysis.io import (
    PredictionFileError,
    parse_prediction_file,
)

_REPO_ROOT = Path(__file__).parent.parent
_DEMO_YAML = _REPO_ROOT / "configs" / "demo.yaml"

_NL = chr(10)


@pytest.fixture
def demo_spec():
    return load_problem(_DEMO_YAML)


@pytest.fixture
def schema_y1(demo_spec):
    # 회귀 대상만 활성화: 입력 x1(float), x2(float), x3(category)
    return build_input_schema(demo_spec, active_target_ids=["synthetic_y1"])


def _write_csv(path: Path, lines: list, bom: bool = True) -> Path:
    enc = "utf-8-sig" if bom else "utf-8"
    path.write_text(_NL.join(lines) + _NL, encoding=enc)
    return path


def _valid_header(schema):
    cols = ["candidate_id"] + list(schema.input_ids) + ["_input_schema_id"]
    return ",".join(cols)


class TestValidFile:
    def test_valid_file_returns_candidates(self, schema_y1, tmp_path):
        sid = schema_y1.schema_id
        p = _write_csv(tmp_path / "ok.csv", [
            "candidate_id,synthetic_x1,synthetic_x2,synthetic_x3,_input_schema_id",
            "c1,1.0,2.0,cat_a," + sid,
            "c2,3.0,4.0,cat_b," + sid,
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["total_error_count"] == 0
        assert res["candidates"] is not None
        assert len(res["candidates"].data) == 2

    def test_column_order_independent(self, schema_y1, tmp_path):
        # 열 순서를 섞어도 ID 로 읽는다
        sid = schema_y1.schema_id
        p = _write_csv(tmp_path / "order.csv", [
            "synthetic_x3,synthetic_x2,synthetic_x1,candidate_id,_input_schema_id",
            "cat_a,2.0,1.0,c1," + sid,
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["total_error_count"] == 0
        assert res["candidates"] is not None
        assert list(res["candidates"].data.columns) == list(schema_y1.input_ids)

    def test_candidate_id_optional(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "nocand.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3",
            "1.0,2.0,cat_a",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["total_error_count"] == 0
        assert res["candidates"].candidate_ids == ["row_1"]


class TestLastRowError:
    def test_last_row_error_fails_whole_file(self, schema_y1, tmp_path):
        # 마지막 행에 오류 (비수치)
        p = _write_csv(tmp_path / "last.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3",
            "1.0,2.0,cat_a",
            "3.0,4.0,cat_b",
            "notnum,5.0,cat_c",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["total_error_count"] >= 1
        assert res["candidates"] is None
        # 마지막 행(3)의 오류가 수집됐는지
        assert any(i.row == 3 and i.code == "NOT_NUMERIC" for i in res["issues"])


class TestErrorDisplayLimit:
    def test_display_limited_but_total_counted(self, schema_y1, tmp_path):
        # 비수치 오류 5건 생성
        lines = ["synthetic_x1,synthetic_x2,synthetic_x3"]
        for i in range(5):
            lines.append("bad,2.0,cat_a")
        p = _write_csv(tmp_path / "many.csv", lines)
        res = parse_prediction_file(p, schema_y1, max_error_display=2)
        assert res["total_error_count"] == 5
        assert len(res["issues"]) == 2
        assert any("5" in n and "2" in n for n in res["notes"])


class TestBadNotTreatedAsMissing:
    def test_nonempty_bad_value_is_error_not_missing(self, schema_y1, tmp_path):
        # 잘못된 문자열이 결측으로 바뀌어 통과하면 안 된다
        p = _write_csv(tmp_path / "bad.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3",
            "abc,2.0,cat_a",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "NOT_NUMERIC" for i in res["issues"])

    def test_non_finite_rejected(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "inf.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3",
            "inf,2.0,cat_a",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "NOT_FINITE" for i in res["issues"])

    def test_unknown_category_rejected(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "cat.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3",
            "1.0,2.0,unknown_cat",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "UNKNOWN_CATEGORY" for i in res["issues"])

    def test_missing_required_rejected(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "miss.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3",
            ",2.0,cat_a",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "MISSING_REQUIRED" for i in res["issues"])


class TestFileLevelErrors:
    def test_duplicate_header_raises(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "dup.csv", [
            "synthetic_x1,synthetic_x1,synthetic_x2,synthetic_x3",
            "1.0,9.0,2.0,cat_a",
        ])
        with pytest.raises(PredictionFileError):
            parse_prediction_file(p, schema_y1)

    def test_empty_file_raises(self, schema_y1, tmp_path):
        p = tmp_path / "empty.csv"
        p.write_text("synthetic_x1,synthetic_x2,synthetic_x3" + _NL, encoding="utf-8-sig")
        with pytest.raises(PredictionFileError):
            parse_prediction_file(p, schema_y1)

    def test_unsupported_format_raises(self, schema_y1, tmp_path):
        p = tmp_path / "x.json"
        p.write_text("{}", encoding="utf-8")
        with pytest.raises(PredictionFileError):
            parse_prediction_file(p, schema_y1)

    def test_mixed_schema_id_raises(self, schema_y1, tmp_path):
        sid = schema_y1.schema_id
        p = _write_csv(tmp_path / "mixed.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3,_input_schema_id",
            "1.0,2.0,cat_a," + sid,
            "3.0,4.0,cat_b,other_schema_id",
        ])
        with pytest.raises(PredictionFileError):
            parse_prediction_file(p, schema_y1)

    def test_wrong_schema_id_raises(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "wrong.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3,_input_schema_id",
            "1.0,2.0,cat_a,deadbeef0000",
        ])
        with pytest.raises(PredictionFileError):
            parse_prediction_file(p, schema_y1)

    def test_unknown_column_is_error(self, schema_y1, tmp_path):
        p = _write_csv(tmp_path / "extra.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3,bogus_col",
            "1.0,2.0,cat_a,xyz",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "UNKNOWN_COLUMN" for i in res["issues"])

    def test_missing_required_column_is_error(self, schema_y1, tmp_path):
        # synthetic_x2 열 자체가 없음
        p = _write_csv(tmp_path / "misscol.csv", [
            "synthetic_x1,synthetic_x3",
            "1.0,cat_a",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "MISSING_COLUMN" for i in res["issues"])


class TestResultColumnDropped:
    def test_target_column_is_unknown_here(self, schema_y1, tmp_path):
        # 예측 파일에 목표 열(synthetic_y1)이 들어오면 미정의 열로 처리된다
        # (입력 규격에 없는 열이므로 UNKNOWN_COLUMN)
        p = _write_csv(tmp_path / "withtarget.csv", [
            "synthetic_x1,synthetic_x2,synthetic_x3,synthetic_y1",
            "1.0,2.0,cat_a,99.0",
        ])
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "UNKNOWN_COLUMN" for i in res["issues"])


class TestExcelPredictionFile:
    """실제 .xlsx 파일로 정상 통과·수식 거부·병합 셀 거부를 검증한다 (작업 9)."""

    def _make_xlsx(self, path, header, rows, formula_cell=None, merge=None):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(header)
        for r in rows:
            ws.append(r)
        if formula_cell is not None:
            cell_ref, formula = formula_cell
            ws[cell_ref] = formula
        if merge is not None:
            ws.merge_cells(merge)
        wb.save(path)
        return path

    def test_valid_xlsx_passes(self, schema_y1, tmp_path):
        sid = schema_y1.schema_id
        p = tmp_path / "ok.xlsx"
        self._make_xlsx(
            p,
            ["synthetic_x1", "synthetic_x2", "synthetic_x3", "_input_schema_id"],
            [[1.0, 2.0, "cat_a", sid], [3.0, 4.0, "cat_b", sid]],
        )
        res = parse_prediction_file(p, schema_y1)
        assert res["total_error_count"] == 0
        assert res["candidates"] is not None
        assert len(res["candidates"].data) == 2

    def test_formula_cell_rejected(self, schema_y1, tmp_path):
        p = tmp_path / "formula.xlsx"
        # 데이터 영역(2행 B열)에 수식을 넣는다
        self._make_xlsx(
            p,
            ["synthetic_x1", "synthetic_x2", "synthetic_x3"],
            [[1.0, 2.0, "cat_a"]],
            formula_cell=("B2", "=1+1"),
        )
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "FORMULA_CELL" for i in res["issues"])

    def test_merged_cell_rejected(self, schema_y1, tmp_path):
        p = tmp_path / "merged.xlsx"
        self._make_xlsx(
            p,
            ["synthetic_x1", "synthetic_x2", "synthetic_x3"],
            [[1.0, 2.0, "cat_a"], [3.0, 4.0, "cat_b"]],
            merge="A2:B2",
        )
        res = parse_prediction_file(p, schema_y1)
        assert res["candidates"] is None
        assert any(i.code == "MERGED_CELL" for i in res["issues"])
